from pathlib import Path
import csv,json,tempfile,zipfile,xml.etree.ElementTree as ET,sys,subprocess,shutil
from common import *

PRESENTATIONS={'.ppt','.pptx','.pptm','.pps','.ppsx','.odp','.key'}
WORDS={'.doc','.docx','.docm','.rtf','.odt','.pages'}
SHEETS={'.xls','.xlsx','.xlsm','.xlsb','.ods','.csv','.tsv','.numbers'}

def family(path):
    ext=Path(path).suffix.lower()
    return 'presentation' if ext in PRESENTATIONS else 'word' if ext in WORDS else 'sheet'

def native(path,fmt,target):
    path=Path(path).resolve()
    target=Path(target).resolve()
    if sys.platform=='darwin':
        # 0. Apple Pages (.pages) -> PDF / DOCX / TXT
        if path.suffix.lower() == '.pages':
            exp_format = 'PDF' if fmt == 'pdf' else 'Microsoft Word' if fmt in ('docx', 'doc') else 'unformatted text' if fmt == 'txt' else None
            if exp_format:
                subprocess.run(['open', '-g', '-j', '-a', 'Pages'], capture_output=True)
                as_code = f'set inPath to POSIX file "{path}"\nset outPath to POSIX file "{target}"\ntell application "Pages"\nset myDoc to open inPath\nexport myDoc to outPath as {exp_format}\nclose myDoc saving no\nend tell'
                try:
                    res = subprocess.run(['osascript', '-e', as_code], capture_output=True, timeout=45)
                    if res.returncode == 0 and target.exists() and target.stat().st_size > 0:
                        return
                except Exception: pass
            # Fallback: Extract QuickLook/Preview.pdf if available
            try:
                if zipfile.is_zipfile(path):
                    with zipfile.ZipFile(path) as z:
                        if 'QuickLook/Preview.pdf' in z.namelist():
                            if fmt == 'pdf':
                                with z.open('QuickLook/Preview.pdf') as src, open(target, 'wb') as dst:
                                    shutil.copyfileobj(src, dst)
                                if target.exists() and target.stat().st_size > 0:
                                    return
                            elif fmt in ('docx', 'doc'):
                                temp_pdf = target.parent / f'{path.stem}_preview.pdf'
                                with z.open('QuickLook/Preview.pdf') as src, open(temp_pdf, 'wb') as dst:
                                    shutil.copyfileobj(src, dst)
                                import documents
                                docx_out = documents.pdf_word(temp_pdf, {})
                                temp_pdf.unlink(missing_ok=True)
                                if Path(docx_out).exists():
                                    shutil.move(docx_out, target)
                                    return
            except Exception: pass

        # Apple Keynote (.key) -> PDF / PPTX
        if path.suffix.lower() == '.key':
            exp_format = 'PDF' if fmt == 'pdf' else 'Microsoft PowerPoint' if fmt in ('pptx', 'ppt') else None
            if exp_format:
                subprocess.run(['open', '-g', '-j', '-a', 'Keynote'], capture_output=True)
                as_code = f'set inPath to POSIX file "{path}"\nset outPath to POSIX file "{target}"\ntell application "Keynote"\nset myDoc to open inPath\nexport myDoc to outPath as {exp_format}\nclose myDoc saving no\nend tell'
                try:
                    res = subprocess.run(['osascript', '-e', as_code], capture_output=True, timeout=45)
                    if res.returncode == 0 and target.exists() and target.stat().st_size > 0:
                        return
                except Exception: pass

        # Apple Numbers (.numbers) -> PDF / XLSX / CSV
        if path.suffix.lower() == '.numbers':
            exp_format = 'PDF' if fmt == 'pdf' else 'Microsoft Excel' if fmt in ('xlsx', 'xls') else 'CSV' if fmt == 'csv' else None
            if exp_format:
                subprocess.run(['open', '-g', '-j', '-a', 'Numbers'], capture_output=True)
                as_code = f'set inPath to POSIX file "{path}"\nset outPath to POSIX file "{target}"\ntell application "Numbers"\nset myDoc to open inPath\nexport myDoc to outPath as {exp_format}\nclose myDoc saving no\nend tell'
                try:
                    res = subprocess.run(['osascript', '-e', as_code], capture_output=True, timeout=45)
                    if res.returncode == 0 and target.exists() and target.stat().st_size > 0:
                        return
                except Exception: pass

        # Native textutil for RTF / RTFD / DOCX / DOC -> HTML / TXT
        if path.suffix.lower() in ('.rtf', '.rtfd', '.docx', '.doc') and fmt in ('txt', 'html'):
            try:
                res = subprocess.run(['textutil', '-convert', fmt, '-output', str(target), str(path)], capture_output=True, timeout=30)
                if res.returncode == 0 and target.exists() and target.stat().st_size > 0:
                    return
            except Exception: pass

        if path.suffix.lower() in ('.rtf', '.rtfd') and fmt in ('docx', 'doc'):
            try:
                res = subprocess.run(['textutil', '-convert', 'docx', '-output', str(target), str(path)], capture_output=True, timeout=30)
                if res.returncode == 0 and target.exists() and target.stat().st_size > 0:
                    return
            except Exception: pass

        # Apple Pages for DOCX / DOC -> PAGES / EPUB
        if path.suffix.lower() in ('.docx', '.doc') and fmt in ('pages', 'epub'):
            subprocess.run(['open', '-g', '-j', '-a', 'Pages'], capture_output=True)
            as_action = f'save front document in (POSIX file "{target}")' if fmt == 'pages' else f'export front document to (POSIX file "{target}") as EPUB'
            as_code = f'tell application "Pages"\nopen (POSIX file "{path}")\ndelay 0.5\n{as_action}\nclose front document saving no\nend tell'
            try:
                res = subprocess.run(['osascript', '-e', as_code], capture_output=True, timeout=45)
                if res.returncode == 0 and target.exists() and target.stat().st_size > 0:
                    return
            except Exception: pass

        # 1. Try LibreOffice if installed
        for soffice in ['/Applications/LibreOffice.app/Contents/MacOS/soffice', 'soffice']:
            if shutil.which(soffice) or Path(soffice).exists():
                try:
                    subprocess.run([soffice, '--headless', '--convert-to', fmt, str(path), '--outdir', str(target.parent)], check=True, timeout=60)
                    out_name = target.parent / (path.stem + '.' + fmt)
                    if out_name.exists():
                        out_name.replace(target)
                        return
                except Exception: pass
        # 2. Try Microsoft Word via AppleScript on macOS
        if path.suffix.lower() in ('.doc','.docx') and fmt=='pdf':
            as_code = f'tell application "Microsoft Word"\nopen "{path}"\nsave as active document file name "{target}" file format format PDF\nclose active document saving no\nend tell'
            try:
                res = subprocess.run(['osascript', '-e', as_code], capture_output=True, timeout=15)
                if res.returncode == 0 and target.exists() and target.stat().st_size > 0:
                    return
            except Exception: pass
        # 3. High-fidelity native macOS paginated PDF converter (doc2pdf-mac)
        if fmt=='pdf' and path.suffix.lower() in ('.doc','.docx','.rtf','.odt','.txt','.html','.htm'):
            bin_candidates = [
                BASE / 'doc2pdf-mac',
                Path(__file__).parent / 'doc2pdf-mac',
                Path('/Applications/DropDial.app/Contents/backend/doc2pdf-mac')
            ]
            for bin_path in bin_candidates:
                if bin_path.exists() and os.access(str(bin_path), os.X_OK):
                    try:
                        res = subprocess.run([str(bin_path), str(path), str(target)], capture_output=True, timeout=30)
                        if res.returncode == 0 and target.exists() and target.stat().st_size > 0:
                            return
                    except Exception: pass
            # Fallback to swift runner
            swift_script_path = Path(__file__).parent / 'doc2pdf-mac.swift'
            if swift_script_path.exists():
                try:
                    res = subprocess.run(['swift', str(swift_script_path), str(path), str(target)], capture_output=True, timeout=30)
                    if res.returncode == 0 and target.exists() and target.stat().st_size > 0:
                        return
                except Exception: pass
        # 4. Spreadsheets to PDF via styled HTML table
        if fmt=='pdf' and path.suffix.lower() in ('.csv','.tsv','.xlsx','.xlsm'):
            try:
                sheets = data_rows(path, target.parent)
                html_parts = ['<!DOCTYPE html><html><head><meta charset="utf-8"><style>',
                              'body { font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif; padding: 20px; font-size: 12px; }',
                              'h2 { font-size: 16px; margin-top: 16px; margin-bottom: 8px; color: #111; }',
                              'table { border-collapse: collapse; width: 100%; margin-bottom: 24px; }',
                              'th, td { border: 1px solid #ddd; padding: 6px 10px; text-align: left; }',
                              'th { background-color: #f5f5f7; font-weight: 600; }',
                              'tr:nth-child(even) { background-color: #fafafa; }',
                              '</style></head><body>']
                for s in sheets:
                    html_parts.append(f'<h2>{s.get("Name", "Sheet")}</h2><table>')
                    for r_idx, row in enumerate(s.get('Rows', [])):
                        html_parts.append('<tr>')
                        tag = 'th' if r_idx == 0 else 'td'
                        for cell in row:
                            val = str(cell) if cell is not None else ''
                            html_parts.append(f'<{tag}>{val}</{tag}>')
                        html_parts.append('</tr>')
                    html_parts.append('</table>')
                html_parts.append('</body></html>')
                temp_html = target.parent / f'{path.stem}_table.html'
                temp_html.write_text(''.join(html_parts), encoding='utf-8')
                bin_candidates = [
                    BASE / 'doc2pdf-mac',
                    Path(__file__).parent / 'doc2pdf-mac',
                    Path('/Applications/DropDial.app/Contents/backend/doc2pdf-mac')
                ]
                for bin_path in bin_candidates:
                    if bin_path.exists() and os.access(str(bin_path), os.X_OK):
                        subprocess.run([str(bin_path), str(temp_html), str(target)], capture_output=True, timeout=30)
                        temp_html.unlink(missing_ok=True)
                        if target.exists() and target.stat().st_size > 0:
                            return
            except Exception: pass
    exe = BASE/'DropDial.exe' if (BASE/'DropDial.exe').exists() else BASE/'DropDial.exe'
    if exe.exists():
        process([exe,'--office-export',path,fmt,target],timeout=300)
    else:
        raise RuntimeError(T('The current system cannot convert this Office format. Please install LibreOffice or Microsoft Word.'))

def data_rows(path,folder):
    ext=Path(path).suffix.lower()
    if ext in ('.csv','.tsv'):
        with open(path,encoding='utf-8-sig',newline='') as stream:return [{'Name':Path(path).stem,'Rows':list(csv.reader(stream,delimiter='\t' if ext=='.tsv' else ','))}]
    if ext in ('.xlsx','.xlsm'):
        try:
            import winreg
            with winreg.OpenKey(winreg.HKEY_CLASSES_ROOT,'Excel.Application\\CLSID'):pass
        except Exception:
            from openpyxl import load_workbook
            workbook=load_workbook(path,data_only=True,read_only=True);result=[{'Name':sheet.title,'Rows':[list(row) for row in sheet.iter_rows(values_only=True)]} for sheet in workbook.worksheets];workbook.close();return result
    output=folder/'sheets.json';native(path,'sheet-data',output);return json.loads(output.read_text(encoding='utf8'))

def text_without_office(path):
    ext=Path(path).suffix.lower()
    if ext in ('.pptx','.pptm','.ppsx'):
        from pptx import Presentation
        document=Presentation(path);parts=[]
        for i,slide in enumerate(document.slides):
            parts.append(T('Page {0}').format(i+1))
            for shape in slide.shapes:
                if shape.has_text_frame:parts.append(shape.text)
                if shape.has_table:
                    parts.extend('\t'.join(cell.text for cell in row.cells) for row in shape.table.rows)
        return '\n\n'.join(parts)
    if ext in ('.docx','.docm'):
        from docx import Document
        document=Document(path);parts=[p.text for p in document.paragraphs]
        parts.extend('\t'.join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
        return '\n'.join(parts)
    raise ValueError(T('This text export needs Microsoft Office'))

def pdf_file(path,folder):
    if Path(path).suffix.lower() in ('.csv','.tsv'):path=csv_workbook(path,folder)
    output=folder/'document.pdf';native(path,'pdf',output);return output

def csv_workbook(path,folder):
    from openpyxl import Workbook
    workbook=Workbook();worksheet=workbook.active
    for row in data_rows(path,folder)[0]['Rows']:worksheet.append(row)
    for row in worksheet:
        for cell in row:
            if isinstance(cell.value,str):cell.data_type='s'
    intermediate=folder/'sheet.xlsx';workbook.save(intermediate);return intermediate

def convert(path,fmt,params):
    path=Path(path);kind=family(path)
    with tempfile.TemporaryDirectory(prefix='dropdial-office-') as temporary:
        folder=Path(temporary)
        if fmt in ('png','jpg'):
            import pypdfium2 as pdfium
            pdf=pdf_file(path,folder);document=pdfium.PdfDocument(pdf)
            if len(document)==1:
                with output_file(path,'converted','.'+fmt) as state:save_image(document[0].render(scale=300/72).to_pil(),state['temp'],fmt,92)
            else:
                with output_folder(path,'slides-'+fmt if kind=='presentation' else 'pages-'+fmt) as state:
                    for index in range(len(document)):save_image(document[index].render(scale=300/72).to_pil(),state['temp']/f'page-{index+1:04d}.{fmt}',fmt,92)
            document.close();return state['output']
        if fmt in ('csv','tsv','json','txt') and kind=='sheet':
            sheets=data_rows(path,folder)
            if fmt=='json':
                with output_file(path,'converted','.json') as state:state['temp'].write_text(json.dumps(sheets,ensure_ascii=False,indent=2),encoding='utf8')
            elif len(sheets)==1:
                with output_file(path,'converted','.'+fmt) as state:
                    with open(state['temp'],'w',encoding='utf-8-sig',newline='') as stream:csv.writer(stream,delimiter=',' if fmt=='csv' else '\t').writerows(sheets[0]['Rows'])
            else:
                with output_folder(path,'sheets-'+fmt) as state:
                    for index,sheet in enumerate(sheets):
                        name=''.join('_' if c in '<>:"/\\|?*' else c for c in sheet['Name']).strip('. ') or f'Sheet{index+1}'
                        with open(state['temp']/f'{index+1:02d}-{name}.{fmt}','w',encoding='utf-8-sig',newline='') as stream:csv.writer(stream,delimiter=',' if fmt=='csv' else '\t').writerows(sheet['Rows'])
            return state['output']
        if fmt=='txt':
            with output_file(path,'converted','.txt') as state:
                if path.suffix.lower() in ('.pptx','.pptm','.ppsx','.docx','.docm'):state['temp'].write_text(text_without_office(path),encoding='utf-8-sig')
                else:native(path,'txt',state['temp'])
            return state['output']
        if fmt=='html':
            with output_folder(path,'html') as state:native(path,'html',state['temp']/'index.html')
            return state['output']
        if path.suffix.lower() in ('.csv','.tsv') and fmt in ('xlsx','xls','ods','pdf'):
            if fmt == 'pdf':
                with output_file(path, 'converted', '.pdf') as state:
                    native(path, 'pdf', state['temp'])
                return state['output']
            intermediate=csv_workbook(path,folder)
            with output_file(path,'converted','.'+fmt) as state:
                if fmt=='xlsx':shutil.copyfile(intermediate,state['temp'])
                else:native(intermediate,fmt,state['temp'])
            return state['output']
        with output_file(path,'converted','.'+fmt) as state:native(path,fmt,state['temp'])
        return state['output']

def metadata(path):
    if not zipfile.is_zipfile(path):return {'Format':Path(path).suffix.upper(),'Details':T('Metadata editing is not supported for this older format')}
    result={}
    with zipfile.ZipFile(path) as package:
        for name in ('docProps/core.xml','docProps/app.xml'):
            if name in package.namelist():
                for node in ET.fromstring(package.read(name)):
                    if node.text:result[node.tag.split('}')[-1]]=node.text
    return result

def edit_metadata(path,params):
    if not zipfile.is_zipfile(path):raise ValueError(T('Metadata editing supports modern Office formats; convert to PPTX/DOCX/XLSX first'))
    values=json.loads(params.get('metadata','{}'));remove=truth(params.get('remove','true'))
    with output_file(path,'metadata',Path(path).suffix.lower()) as state:
        with zipfile.ZipFile(path) as source,zipfile.ZipFile(state['temp'],'w',zipfile.ZIP_DEFLATED) as target:
            for entry in source.infolist():
                data=source.read(entry.filename)
                if entry.filename in ('docProps/core.xml','docProps/app.xml'):
                    root=ET.fromstring(data)
                    for node in list(root):
                        key=node.tag.split('}')[-1]
                        if remove and key not in ('Application','AppVersion','Slides','Words','Pages','Characters','Paragraphs','ScaleCrop','SharedDoc','LinksUpToDate','HyperlinksChanged'):node.text=''
                        elif not remove and key in values:node.text='' if values[key] is None else str(values[key])
                    data=ET.tostring(root,encoding='utf8',xml_declaration=True)
                if remove and entry.filename=='docProps/custom.xml':
                    root=ET.fromstring(data)
                    for node in list(root):root.remove(node)
                    data=ET.tostring(root,encoding='utf8',xml_declaration=True)
                target.writestr(entry,data)
    return state['output']

def inspect(path,folder):
    import pypdfium2 as pdfium
    folder=Path(folder);pdf=pdf_file(path,folder);document=pdfium.PdfDocument(pdf);preview=folder/'preview.png';document[0].render(scale=1.3).to_pil().save(preview)
    value={'Kind':'office','Preview':str(preview),'Width':0,'Height':0,'Duration':0,'FrameRate':0,'Pages':len(document),'Metadata':metadata(path)};document.close();return value

def selected_pdf(path,params):
    from pypdf import PdfReader,PdfWriter
    with tempfile.TemporaryDirectory(prefix='dropdial-office-') as temporary:
        source=pdf_file(path,Path(temporary));reader=PdfReader(source);order=params.get('pageOrder','').strip();indices=[int(v.strip())-1 for v in order.split(',')] if order else list(range(len(reader.pages)));writer=PdfWriter()
        for index in indices:
            if index<0 or index>=len(reader.pages):raise ValueError(T('A page number is outside the document'))
            writer.add_page(reader.pages[index])
        with output_file(path,'selected','.pdf') as state:
            with open(state['temp'],'wb') as stream:writer.write(stream)
        writer.close();return state['output']

def merge(paths,params):
    from pypdf import PdfReader,PdfWriter
    writer=PdfWriter()
    with tempfile.TemporaryDirectory(prefix='dropdial-office-merge-') as temporary:
        for index,path in enumerate(paths):
            folder=Path(temporary)/str(index);folder.mkdir();pdf=pdf_file(path,folder);writer.append(PdfReader(pdf))
        with output_file(paths[0],'merged','.pdf') as state:
            with open(state['temp'],'wb') as stream:writer.write(stream)
    writer.close();return state['output']
